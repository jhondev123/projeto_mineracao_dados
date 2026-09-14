"""Gera uma planilha Excel com os dados exatamente como vieram das APIs, antes
de qualquer tratamento: uma aba por fonte, mais a aba "Leia-me".

Serve para mostrar na apresentacao o ponto de partida (o que foi coletado).
O dataset tratado continua sendo gerado pelo dataset.py, com as mesmas fontes.

Uso:
    python planilha_bruta.py
    python planilha_bruta.py --periodos 201501-202612
"""

import argparse
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import sidra
from dataset import BCB_SERIES, IPCA_TABELA, IPCA_VARIAVEL, PASTA_DADOS, SETORES, buscar_bcb

LOCALIDADES = "N1[all]|N3[all]"  # Brasil e UFs, igual ao dataset.py

FONTE = "Arial"
FUNDO_CABECALHO = PatternFill("solid", fgColor="E1E0D9")

ROTULOS_SETOR = {
    "COMERCIO_VAREJISTA": (
        "varejo",
        "Pesquisa Mensal de Comércio (PMC)",
        "Índice de receita nominal de vendas do comércio varejista (média de 2022 = 100), sem ajuste sazonal",
    ),
    "SERVICOS": (
        "serviços",
        "Pesquisa Mensal de Serviços (PMS)",
        "Índice de receita nominal de serviços (média de 2022 = 100), sem ajuste sazonal",
    ),
}

ROTULOS_BCB = {
    "dolar": ("dólar", "Taxa de câmbio livre, dólar americano (venda), média do mês, em R$ por US$"),
    "selic": ("Selic", "Taxa Selic acumulada no mês, anualizada, em % ao ano"),
}

LARGURAS = {
    "Id da variável": 13,
    "Variável": 40,
    "Unidade": 15,
    "Classificação": 16,
    "Categoria": 55,
    "Nível territorial": 22,
    "Código da localidade": 12,
    "Localidade": 22,
    "Período": 10,
    "Data": 12,
    "Valor": 12,
}

COLUNAS_IBGE = [
    ("Id da variável", "código da variável na tabela", "id"),
    ("Variável", "nome da variável", "variavel"),
    ("Unidade", "unidade de medida", "unidade"),
    ("Classificação", "quebra usada na tabela (ex.: Tipos de índice)", "classificacoes → nome"),
    ("Categoria", "valor dessa quebra (ex.: índice de receita nominal)", "classificacoes → categoria"),
    ("Nível territorial", "Brasil ou Unidade da Federação", "localidade → nivel → nome"),
    ("Código da localidade", "código do IBGE (1 = Brasil, 35 = São Paulo)", "localidade → id"),
    ("Localidade", "nome do Brasil ou do estado", "localidade → nome"),
    ("Período", "ano e mês no formato AAAAMM (202504 = abril de 2025)", "chave de serie"),
    ("Valor", "valor publicado", "valor de serie"),
]

COLUNAS_BCB = [
    ("Data", "dia/mês/ano; as séries mensais vêm datadas no dia 1", "data"),
    ("Valor", "valor publicado", "valor"),
]

SIMBOLOS_IBGE = [
    ("-", "zero absoluto (não houve)"),
    ("0", "zero por arredondamento"),
    ("X", "valor omitido por sigilo estatístico"),
    ("..", "não se aplica"),
    ("...", "dado não disponível"),
]


def valor_celula(valor):
    """Numero vira numero no Excel; simbolo do IBGE (X, .., ...) fica como texto."""
    try:
        return float(valor)
    except (TypeError, ValueError):
        return valor


def linhas_ibge(payload, com_classificacao):
    """Achata o JSON do IBGE sem alterar nada: uma linha por localidade x periodo."""
    for variavel in payload:
        for resultado in variavel["resultados"]:
            classificacoes = resultado.get("classificacoes", [])
            nomes = "; ".join(c["nome"] for c in classificacoes)
            categorias = "; ".join(nome for c in classificacoes for nome in c["categoria"].values())
            for serie in resultado["series"]:
                localidade = serie["localidade"]
                for periodo, valor in serie["serie"].items():
                    linha = [variavel["id"], variavel["variavel"], variavel["unidade"]]
                    if com_classificacao:
                        linha += [nomes, categorias]
                    linha += [
                        localidade.get("nivel", {}).get("nome"),
                        localidade["id"],
                        localidade["nome"],
                        periodo,
                        valor_celula(valor),
                    ]
                    yield linha


def escrever_aba(wb, nome, cabecalho, linhas):
    aba = wb.create_sheet(nome)
    aba.append(cabecalho)
    for linha in linhas:
        aba.append(linha)

    fonte = Font(name=FONTE, size=10)
    for linha in aba.iter_rows():
        for celula in linha:
            celula.font = fonte
    for celula in aba[1]:
        celula.font = Font(name=FONTE, size=10, bold=True)
        celula.fill = FUNDO_CABECALHO

    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    for i, titulo in enumerate(cabecalho, start=1):
        aba.column_dimensions[get_column_letter(i)].width = LARGURAS[titulo]
    return aba.max_row - 1


def escrever_leia_me(aba, abas, coleta):
    negrito = Font(name=FONTE, size=10, bold=True)
    normal = Font(name=FONTE, size=10)
    quebra = Alignment(wrap_text=True, vertical="top")

    def linha(valores, fonte=normal, cabecalho=False):
        aba.append(valores)
        for celula in aba[aba.max_row]:
            celula.font = fonte
            celula.alignment = quebra
            if cabecalho:
                celula.fill = FUNDO_CABECALHO

    aba.append(["Dados brutos coletados"])
    aba["A1"].font = Font(name=FONTE, size=14, bold=True)
    linha(["Faturamento do comércio varejista e dos serviços por estado — projeto de mineração de dados (CRISP-DM)"])
    linha([f"Coletados em {coleta:%d/%m/%Y às %H:%M}, exatamente como as APIs devolvem, antes de qualquer tratamento."])
    linha([])

    linha(["Aba", "Fonte", "O que é", "Consulta feita"], negrito, cabecalho=True)
    for valores in abas:
        linha(list(valores))
    linha([])

    linha(["Colunas das abas do IBGE", "Significado", "Campo no JSON da API"], negrito, cabecalho=True)
    for valores in COLUNAS_IBGE:
        linha(list(valores))
    linha(["", "A aba do IPCA não tem Classificação nem Categoria: a tabela 1737 não tem quebras."])
    linha([])

    linha(["Colunas das abas do Banco Central", "Significado", "Campo no JSON da API"], negrito, cabecalho=True)
    for valores in COLUNAS_BCB:
        linha(list(valores))
    linha([])

    linha(["Símbolo do IBGE", "Significado"], negrito, cabecalho=True)
    for valores in SIMBOLOS_IBGE:
        linha(list(valores))
    linha([])

    linha(["Observações"], negrito)
    linha(["Nenhuma linha foi removida, juntada, convertida ou renomeada. Os números foram gravados como número "
           "para o Excel conseguir ler; símbolos do IBGE, se aparecerem, ficam como texto."])
    linha(["O dataset tratado (CSV) é gerado pelo dataset.py a partir destas mesmas consultas."])

    for coluna, largura in zip("ABCD", (26, 60, 70, 60)):
        aba.column_dimensions[coluna].width = largura


def main():
    parser = argparse.ArgumentParser(description="Planilha com os dados brutos das APIs")
    parser.add_argument("--periodos", default="201201-202612", help="intervalo AAAAMM-AAAAMM")
    args = parser.parse_args()

    coleta = datetime.now()
    wb = Workbook()
    leia_me = wb.active
    leia_me.title = "Leia-me"
    abas = []

    for setor, (tabela, variavel, filtro) in SETORES.items():
        apelido, pesquisa, descricao = ROTULOS_SETOR[setor]
        nome = f"IBGE {apelido} {tabela}"
        print(f"[{nome}] baixando...", flush=True)
        payload = sidra.buscar(tabela, periodos=args.periodos, localidades=LOCALIDADES,
                               classificacoes=[filtro], variaveis=variavel)
        cabecalho = [titulo for titulo, _, _ in COLUNAS_IBGE]
        total = escrever_aba(wb, nome, cabecalho, linhas_ibge(payload, com_classificacao=True))
        abas.append((nome, f"IBGE — {pesquisa}, tabela {tabela}", descricao,
                     f"variável {variavel}, filtro {filtro}, localidades {LOCALIDADES} (Brasil e UFs), "
                     f"períodos {args.periodos}"))
        print(f"[{nome}] {total} linhas")

    nome = f"IBGE IPCA {IPCA_TABELA}"
    print(f"[{nome}] baixando...", flush=True)
    payload = sidra.buscar(IPCA_TABELA, periodos=args.periodos, localidades="N1[all]",
                           classificacoes=[], variaveis=IPCA_VARIAVEL)
    cabecalho = [titulo for titulo, _, _ in COLUNAS_IBGE if titulo not in ("Classificação", "Categoria")]
    total = escrever_aba(wb, nome, cabecalho, linhas_ibge(payload, com_classificacao=False))
    abas.append((nome, f"IBGE — IPCA, tabela {IPCA_TABELA}", "Número-índice do IPCA (dezembro de 1993 = 100)",
                 f"variável {IPCA_VARIAVEL}, localidade N1[all] (Brasil), períodos {args.periodos}"))
    print(f"[{nome}] {total} linhas")

    inicio, fim = args.periodos.split("-")
    for coluna, serie in BCB_SERIES.items():
        apelido, descricao = ROTULOS_BCB[coluna]
        nome = f"BCB {apelido} {serie}"
        print(f"[{nome}] baixando...", flush=True)
        dados = buscar_bcb(serie, args.periodos)
        linhas = ([item["data"], valor_celula(item["valor"])] for item in dados)
        total = escrever_aba(wb, nome, ["Data", "Valor"], linhas)
        abas.append((nome, f"Banco Central — SGS, série {serie}", descricao,
                     f"série {serie}, dataInicial 01/{inicio[4:]}/{inicio[:4]}, dataFinal 01/{fim[4:]}/{fim[:4]}"))
        print(f"[{nome}] {total} linhas")

    escrever_leia_me(leia_me, abas, coleta)

    PASTA_DADOS.mkdir(exist_ok=True)
    caminho = PASTA_DADOS / f"dados_brutos_{coleta:%Y%m%d}.xlsx"
    wb.save(caminho)
    print(f"salvo: {caminho}")


if __name__ == "__main__":
    main()
